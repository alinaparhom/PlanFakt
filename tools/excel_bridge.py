"""Настраиваемый импорт XLSX/XLSB и выгрузка план-факта."""

import calendar
import json
import re
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook


def text(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).lower().replace("ё", "е")


def number(value, default=0):
    if value in (None, ""):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(" ", "").replace("\u00a0", "").replace(",", "."))
    except ValueError:
        return default


def column_index(value, required=False):
    raw = str(value or "").strip().upper()
    if not raw:
        if required:
            raise ValueError("Не заполнена обязательная колонка")
        return None
    if raw.isdigit():
        return int(raw) - 1
    if not re.fullmatch(r"[A-Z]{1,3}", raw):
        raise ValueError(f"Некорректная колонка: {raw}")
    result = 0
    for char in raw:
        result = result * 26 + ord(char) - 64
    return result - 1


def value_at(row, index):
    return row[index] if index is not None and index < len(row) else None


def read_book(source):
    if Path(source).suffix.lower() == ".xlsb":
        try:
            from pyxlsb import open_workbook
        except ImportError as exc:
            raise ValueError("Для XLSB на сервере не установлен модуль pyxlsb") from exc
        book = open_workbook(source)
        try:
            sheets = {}
            for name in book.sheets:
                with book.get_sheet(name) as sheet:
                    sheets[name] = [[item.v for item in row] for row in sheet.rows()]
            return sheets
        finally:
            book.close()
    book = load_workbook(source, data_only=True, read_only=True)
    try:
        return {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)] for sheet in book.worksheets}
    finally:
        book.close()


def find_layout(rows):
    aliases = {
        "name": ("наименование работ", "наименование работы", "работа", "вид работ"),
        "organization": ("подрядчик", "организация", "исполнитель"),
        "code": ("код работы", "шифр", "id работы"),
        "unit": ("ед. изм", "единица измерения", "единица", "изм."),
        "total": ("объем проекта", "общий объем", "объём проекта", "всего объем"),
        "remaining": ("остаток", "оставшийся объем", "оставшийся объём"),
    }
    best = None
    for row_index, row in enumerate(rows[:30]):
        found = {}
        for col_index, value in enumerate(row[:250]):
            normalized = text(value)
            for key, variants in aliases.items():
                if key not in found and any(variant in normalized for variant in variants):
                    found[key] = col_index
        score = int("name" in found) * 3 + int("organization" in found) * 3 + len(found)
        if best is None or score > best[0]:
            best = (score, row_index, found)
    return best[1:] if best and "name" in best[2] and "organization" in best[2] else None


def import_book(source, output, month):
    sheets = read_book(source)
    result = {"month": month, "source": Path(source).name, "sheets": [], "works": [], "warnings": []}
    for sheet_name, rows in sheets.items():
        layout = find_layout(rows)
        if not layout:
            continue
        header_row, columns = layout
        result["sheets"].append(sheet_name)
        for row_index, row in enumerate(rows[header_row + 1:], start=header_row + 2):
            name = str(value_at(row, columns["name"]) or "").strip()
            organization = str(value_at(row, columns["organization"]) or "").strip()
            if not name or not organization:
                continue
            total = number(value_at(row, columns.get("total")))
            result["works"].append({"sourceKey": f"{sheet_name}:{row_index}", "sheet": sheet_name, "row": row_index, "code": str(value_at(row, columns.get("code")) or f"{sheet_name}-{row_index}").strip(), "name": name, "organization": organization, "unit": str(value_at(row, columns.get("unit")) or "ед.").strip(), "totalVolume": total, "priorActual": 0, "sourceRemaining": number(value_at(row, columns.get("remaining")), total), "monthlyPlan": 0, "plans": {}, "factColumns": {}})
    if not result["works"]:
        raise ValueError("Не найдены строки работ. Используйте ручную настройку колонок.")
    Path(output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


def import_mapped(source, output, month, mapping_file):
    mapping = json.loads(Path(mapping_file).read_text(encoding="utf-8"))
    sheets = read_book(source)
    sheet_name = str(mapping.get("sheet") or next(iter(sheets), ""))
    if sheet_name not in sheets:
        raise ValueError(f"Лист «{sheet_name}» не найден")
    rows = sheets[sheet_name]
    required = ("name", "organization", "firstDay")
    cols = {key: column_index(value, key in required) for key, value in mapping.get("columns", {}).items()}
    hierarchy_columns = [column_index(value, True) for value in mapping.get("hierarchyColumns", []) if str(value).strip()]
    start_row = max(1, int(mapping.get("dataStartRow", 2)))
    row_type = cols.get("rowType")
    plan_marker = text(mapping.get("planRowValue") or "План")
    fact_offset = int(mapping.get("factRowOffset", 1))
    days = calendar.monthrange(int(month[:4]), int(month[5:]))[1]
    first_day = cols["firstDay"]
    result = {"month": month, "source": Path(source).name, "sheets": [sheet_name], "mapping": mapping, "works": [], "warnings": []}
    blank_rows = 0
    for source_row, row in enumerate(rows[start_row - 1:], start=start_row):
        name = str(value_at(row, cols.get("name")) or "").strip()
        if row_type is not None and text(value_at(row, row_type)) != plan_marker:
            continue
        if not name:
            blank_rows += 1
            if blank_rows >= 25:
                break
            continue
        blank_rows = 0
        organization = str(value_at(row, cols.get("organization")) or "").strip()
        if not organization:
            result["warnings"].append(f"Строка {source_row}: нет организации")
            continue
        total = number(value_at(row, cols.get("total")))
        prior = number(value_at(row, cols.get("priorActual")))
        remaining = number(value_at(row, cols.get("remaining")), max(0, total - prior))
        hierarchy = [str(value_at(row, index) or "").strip() for index in hierarchy_columns]
        hierarchy = [item for item in hierarchy if item and item != "-"]
        plans = {f"{month}-{day:02d}": number(value_at(row, first_day + day - 1)) for day in range(1, days + 1)}
        result["works"].append({"sourceKey": f"{sheet_name}:{source_row}", "sheet": sheet_name, "row": source_row, "factRow": source_row + fact_offset, "code": str(value_at(row, cols.get("code")) or value_at(row, cols.get("number")) or f"{sheet_name}-{source_row}").strip(), "number": str(value_at(row, cols.get("number")) or "").strip(), "hierarchy": hierarchy, "name": name, "organization": organization, "unit": str(value_at(row, cols.get("unit")) or "ед.").strip(), "totalVolume": total, "priorActual": prior, "sourceRemaining": remaining, "monthlyPlan": number(value_at(row, cols.get("monthlyPlan"))), "plans": plans, "factColumns": {f"{month}-{day:02d}": first_day + day for day in range(1, days + 1)}})
    if not result["works"]:
        raise ValueError("По выбранной схеме не найдено ни одной строки работ")
    Path(output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


def export_book(source, output, facts_file):
    payload = json.loads(Path(facts_file).read_text(encoding="utf-8"))
    if Path(source).suffix.lower() == ".xlsx":
        shutil.copy2(source, output)
        workbook = load_workbook(output, data_only=False, read_only=False)
        for fact in payload.get("facts", []):
            sheet = workbook[fact["sheet"]]
            target = sheet.cell(int(fact["row"]), int(fact["column"]))
            target.value = number(target.value) + number(fact["amount"])
        workbook.save(output)
        return
    parsed = payload.get("parsed") or {}
    fact_values = {}
    for fact in payload.get("facts", []):
        key = (fact.get("sourceKey"), fact.get("date"))
        fact_values[key] = fact_values.get(key, 0) + number(fact.get("amount"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "План-Факт"
    month = parsed.get("month", "")
    days = calendar.monthrange(int(month[:4]), int(month[5:]))[1]
    max_levels = max((len(work.get("hierarchy", [])) for work in parsed.get("works", [])), default=0)
    prefix_headers = ["№", "Код"] + [f"Уровень {i + 1}" for i in range(max_levels)] + ["Наименование работы", "Организация", "Ед. изм.", "Всего по проекту", "Выполнено до месяца", "Остаток", "План на месяц"]
    sheet.append(prefix_headers + ["План/Факт"] + list(range(1, days + 1)))
    for work in parsed.get("works", []):
        hierarchy = list(work.get("hierarchy", [])) + [""] * (max_levels - len(work.get("hierarchy", [])))
        prefix = [work.get("number", ""), work.get("code", "")] + hierarchy + [work.get("name", ""), work.get("organization", ""), work.get("unit", ""), work.get("totalVolume", 0), work.get("priorActual", 0), work.get("sourceRemaining", 0), work.get("monthlyPlan", 0)]
        sheet.append(prefix + ["План"] + [work.get("plans", {}).get(f"{month}-{day:02d}", 0) for day in range(1, days + 1)])
        sheet.append([""] * len(prefix) + ["Факт"] + [fact_values.get((work.get("sourceKey"), f"{month}-{day:02d}"), 0) for day in range(1, days + 1)])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for item in sheet[1]:
        item.font = item.font.copy(bold=True, color="FFFFFF")
        item.fill = item.fill.copy(fill_type="solid", fgColor="246BFE")
    for index in range(1, sheet.max_column + 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 18 if index > 2 else 12
    workbook.save(output)


def main():
    action = sys.argv[1] if len(sys.argv) > 1 else ""
    if action == "import":
        import_book(sys.argv[2], sys.argv[3], sys.argv[4])
    elif action == "import-mapped":
        import_mapped(sys.argv[2], sys.argv[3], sys.argv[4], sys.argv[5])
    elif action == "export":
        export_book(sys.argv[2], sys.argv[3], sys.argv[4])
    else:
        raise ValueError("Неизвестное действие")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)
